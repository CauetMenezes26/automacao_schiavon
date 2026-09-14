"""Testes de `crawler/reports/painel_excel.py::gerar_relatorio`.

Roda contra um arquivo temporario, sem banco — os `dados` sao dicts fake no
mesmo formato que `painel_flow._coletar_dados` monta a partir dos `fetch_*`.

    python -m pytest tests/test_painel_excel.py -v
    python tests/test_painel_excel.py          (sem pytest instalado)
"""

from __future__ import annotations

import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

from crawler.reports.painel_excel import (  # noqa: E402
    ABA_COMPARACAO,
    ABA_DIVERGENCIA,
    gerar_relatorio,
)

_INICIO, _FIM = date(2026, 9, 7), date(2026, 9, 13)
_GERADO_EM = datetime(2026, 9, 9, 14, 0)


def _dados_fake(**overrides) -> dict:
    base = {
        "comparacao_precos": [
            {"arquivo": "wind_CHENEY_09072026.pdf", "fornecedor": "Cheney Brothers",
             "item": "File Mignon", "item_cotado": "FILET MIGNON",
             "qtd": 4, "preco_invoice": 42.90, "preco_referencia": 42.90,
             "dif_unitaria": 0.0, "dif_pct": 0.0, "dif_valor": 0.0, "cod_status": 0},
        ],
        "divergencias": [],
    }
    base.update(overrides)
    return base


def _celulas(ws):
    return [c.value for row in ws.iter_rows() for c in row if c.value is not None]


def test_gerar_relatorio_cria_as_duas_abas():
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(_dados_fake(), caminho, _GERADO_EM, _INICIO, _FIM)

    wb = openpyxl.load_workbook(caminho)
    assert wb.sheetnames == [ABA_COMPARACAO, ABA_DIVERGENCIA]


def test_nome_da_aba_cabe_no_limite_do_excel():
    assert len(ABA_COMPARACAO) <= 31
    assert len(ABA_DIVERGENCIA) <= 31


def test_titulo_na_primeira_celula_de_cada_aba():
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(_dados_fake(), caminho, _GERADO_EM, _INICIO, _FIM)

    wb = openpyxl.load_workbook(caminho)
    assert wb[ABA_COMPARACAO]["A1"].value == "RPA Schiavon - Cotacao x Invoice"
    assert wb[ABA_DIVERGENCIA]["A1"].value == "RPA Schiavon - Divergencias Cotacao x Invoice"


def test_aba_comparacao_traz_arquivo_fornecedor_e_os_dois_precos():
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(
        _dados_fake(comparacao_precos=[
            {"arquivo": "wind_USFOODS_02-09.pdf", "fornecedor": "US Foods",
             "item": "Peito de Frango", "item_cotado": "CHICKEN BREAST",
             "qtd": 10, "preco_invoice": 12.20, "preco_referencia": 12.20,
             "dif_unitaria": 0.0, "dif_pct": 0.0, "dif_valor": 0.0, "cod_status": 0},
        ]),
        caminho, _GERADO_EM, _INICIO, _FIM,
    )

    valores = _celulas(openpyxl.load_workbook(caminho)[ABA_COMPARACAO])
    assert "wind_USFOODS_02-09.pdf" in valores
    assert "US Foods" in valores
    assert 12.20 in valores


def test_aba_comparacao_marca_o_item_como_conciliado():
    """O cliente le "Conciliado", nao "Conferido" — a aba so tem o que fechou."""
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(_dados_fake(), caminho, _GERADO_EM, _INICIO, _FIM)

    valores = _celulas(openpyxl.load_workbook(caminho)[ABA_COMPARACAO])
    assert "Conciliado" in valores
    assert "Conferido" not in valores


def test_aba_comparacao_sem_dados_mostra_nota():
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(_dados_fake(comparacao_precos=[]), caminho, _GERADO_EM, _INICIO, _FIM)

    valores = _celulas(openpyxl.load_workbook(caminho)[ABA_COMPARACAO])
    assert any("Nenhum item conciliado" in str(v) for v in valores)


def test_aba_divergencia_lista_preco_fora_da_tolerancia():
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(
        _dados_fake(divergencias=[
            {"arquivo": "wind_USFOODS_02-09.pdf", "fornecedor": "US Foods",
             "item": "Peito de Frango", "item_cotado": "CHICKEN BREAST",
             "qtd": 10, "preco_invoice": 12.80, "preco_referencia": 12.20,
             "dif_unitaria": 0.60, "dif_pct": 4.9, "dif_valor": 6.00, "cod_status": 10},
        ]),
        caminho, _GERADO_EM, _INICIO, _FIM,
    )

    valores = _celulas(openpyxl.load_workbook(caminho)[ABA_DIVERGENCIA])
    assert "CHICKEN BREAST" in valores
    assert 0.60 in valores
    assert 6.00 in valores
    assert "Acima do cotado" in valores


def test_aba_divergencia_inclui_item_sem_comparacao():
    """Item sem par na cotacao entra na aba de divergencia com preco cotado
    vazio — e o motivo diz por que nao deu para comparar."""
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(
        _dados_fake(divergencias=[
            {"arquivo": "wind_SYSCO_03-09.pdf", "fornecedor": "Sysco",
             "item": "Molho Especial", "item_cotado": None,
             "qtd": 2, "preco_invoice": 8.10, "preco_referencia": None,
             "dif_unitaria": None, "dif_pct": None, "dif_valor": None, "cod_status": 20},
        ]),
        caminho, _GERADO_EM, _INICIO, _FIM,
    )

    ws = openpyxl.load_workbook(caminho)[ABA_DIVERGENCIA]
    valores = _celulas(ws)
    assert "Molho Especial" in valores
    assert "Item nao achado na cotacao" in valores
    assert ws.cell(row=6, column=7).value is None  # preco cotado vazio


def test_aba_divergencia_fecha_com_contagem_e_soma():
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(
        _dados_fake(divergencias=[
            {"arquivo": "a.pdf", "fornecedor": "US Foods", "item": "X", "item_cotado": "X",
             "qtd": 10, "preco_invoice": 12.80, "preco_referencia": 12.20,
             "dif_unitaria": 0.60, "dif_pct": 4.9, "dif_valor": 6.00, "cod_status": 10},
            {"arquivo": "b.pdf", "fornecedor": "Sysco", "item": "Y", "item_cotado": "Y",
             "qtd": 3, "preco_invoice": 5.00, "preco_referencia": 5.50,
             "dif_unitaria": -0.50, "dif_pct": -9.1, "dif_valor": -1.50, "cod_status": 11},
        ]),
        caminho, _GERADO_EM, _INICIO, _FIM,
    )

    valores = _celulas(openpyxl.load_workbook(caminho)[ABA_DIVERGENCIA])
    assert "2 divergencia(s)" in valores
    assert "=SUM(J6:J7)" in valores


def test_aba_divergencia_sem_dados_mostra_nota():
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(_dados_fake(), caminho, _GERADO_EM, _INICIO, _FIM)

    valores = _celulas(openpyxl.load_workbook(caminho)[ABA_DIVERGENCIA])
    assert any("Nenhuma divergencia" in str(v) for v in valores)


def test_sobrescreve_em_vez_de_acumular():
    caminho = Path(tempfile.mkdtemp()) / "painel_operacao.xlsx"
    gerar_relatorio(_dados_fake(), caminho, _GERADO_EM, _INICIO, _FIM)
    gerar_relatorio(_dados_fake(), caminho, _GERADO_EM, _INICIO, _FIM)

    ws = openpyxl.load_workbook(caminho)[ABA_COMPARACAO]
    ocorrencias = [v for v in _celulas(ws) if v == "Cheney Brothers"]
    assert len(ocorrencias) == 1


if __name__ == "__main__":
    falhas = 0
    testes = [(n, o) for n, o in sorted(globals().items())
              if n.startswith("test_") and callable(o)]
    for nome, func in testes:
        try:
            func()
            print(f"  ok   {nome}")
        except AssertionError as exc:
            falhas += 1
            print(f"  FALHA {nome}: {exc or 'assert'}")
        except Exception as exc:  # noqa: BLE001
            falhas += 1
            print(f"  ERRO  {nome}: {type(exc).__name__}: {exc}")
    print(f"\n{len(testes) - falhas}/{len(testes)} passaram")
    sys.exit(1 if falhas else 0)
