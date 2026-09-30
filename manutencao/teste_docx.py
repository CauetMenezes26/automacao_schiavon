from decimal import Decimal
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from commons.matcher import POLine
from commons.paths import REPORTS_DIR
from conciliacao.reconcile_erp import (
    QTY_MISMATCH_PO,
    PRICE_MISMATCH_PO,
    reconcile_items_against_po,
)

po_lines = [POLine(key=1, supplier_unit_id="40102", scancode=None,
                    item_name="Beef Sirloin Top Butt",
                    ordered=Decimal("3"), received=Decimal("3"),
                    invoiced_total_cost=Decimal("1256.09"))]

items = [{"id": 100, "description": "BEEF - Sirloin Top Butt", "quantity": Decimal("10"),
          "unit_price": Decimal("125.60"), "total_price": Decimal("1256.00"),
          "item_code": "40102"}]

resultado = reconcile_items_against_po(items, po_lines)

print(resultado)

# 1. Cabeçalho fake da invoice — no fluxo real isso vem de
#    fetch_invoice_headers_for_reconciliation() (domain/service/conciliacao_service.py).
header = {
    "invoice_number": "INV-2026-001",
    "supplier_name": "Fake Supplier LLC",
    "invoice_date": "2026-09-25",
}


# 2. Separa por tipo de divergência — um item pode entrar nas duas listas
#    se tiver os dois issue_codes ao mesmo tempo.
divergentes_qtd = [item for item in resultado["items"] if QTY_MISMATCH_PO in item["issue_codes"]]
divergentes_valor = [item for item in resultado["items"] if PRICE_MISMATCH_PO in item["issue_codes"]]


def _adicionar_secao(document, titulo, itens, mensagem_vazio):
    document.add_heading(titulo, level=1)
    if not itens:
        document.add_paragraph(mensagem_vazio)
        return
    tabela = document.add_table(rows=1, cols=2)
    tabela.style = "Table Grid"
    cabecalho = tabela.rows[0].cells
    cabecalho[0].text = "Item"
    cabecalho[1].text = "Divergência"
    for item in itens:
        linha = tabela.add_row().cells
        linha[0].text = item["description_invoice"] or ""
        linha[1].text = ", ".join(item["issue_codes"])


# 3. Documento.
document = Document()

PRETO = RGBColor(0x00, 0x00, 0x00)
DOURADO = RGBColor(0xC9, 0xA2, 0x27)


def _forcar_fonte(estilo, cor=None):
    """Muda a fonte de um style inteiro (Normal, Title, Heading 1...) para
    Times New Roman. `rFonts.eastAsia` precisa ser setado à parte — sem isso
    o Word às vezes ignora `font.name` e usa a fonte padrão do tema mesmo
    assim (bug conhecido do python-docx)."""
    estilo.font.name = "Times New Roman"
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if cor is not None:
        estilo.font.color.rgb = cor


_forcar_fonte(document.styles["Normal"])
_forcar_fonte(document.styles["Title"], cor=PRETO)
_forcar_fonte(document.styles["Heading 1"], cor=DOURADO)

document.add_heading(f"Relatório de divergências — Invoice {header['invoice_number']}", 0)
document.add_paragraph(f"Fornecedor: {header['supplier_name']}")
document.add_paragraph(f"Data: {header['invoice_date']}")

_adicionar_secao(document, "Divergência de quantidade", divergentes_qtd,
                  "Nenhuma divergência de quantidade nesta invoice.")
_adicionar_secao(document, "Divergência de valor", divergentes_valor,
                  "Nenhuma divergência de valor nesta invoice.")

nome_arquivo = f"divergencia_{header['invoice_number']}.docx"
document.save(REPORTS_DIR / nome_arquivo)
print(f"Salvo em: {REPORTS_DIR / nome_arquivo}")