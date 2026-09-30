"""Geração de relatório em Word (.docx) — só formatação, sem regra de negócio.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

from commons.logging_config import get_logger

__all__ = ["gerar_relatorio", "gerar_com_seguranca", "DocxReportError"]

log = get_logger(__name__)

_FONTE = "Times New Roman"
_PRETO = RGBColor(0x00, 0x00, 0x00)
_DOURADO = RGBColor(0xC9, 0xA2, 0x27)


class DocxReportError(RuntimeError):
    """Falha ao montar ou salvar o .docx (caminho inválido, disco sem espaço,
    arquivo aberto no Word travando a escrita)."""


# Caracteres de controle que o XML do Word não aceita (python-docx levanta
# ValueError). Vêm de texto lido pelo Vision.
_XML_INVALIDO = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _limpar(texto: str) -> str:
    return _XML_INVALIDO.sub("", texto)


def _forcar_fonte(estilo, cor: RGBColor | None = None) -> None:
    """`rFonts.eastAsia` precisa ser setado à parte de `font.name` — sem
    isso o Word ignora a fonte em alguns runs e cai no padrão do tema
    (bug conhecido do python-docx, confirmado neste projeto)."""
    estilo.font.name = _FONTE
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), _FONTE)
    if cor is not None:
        estilo.font.color.rgb = cor


def gerar_relatorio(
    caminho: Path,
    titulo: str,
    linhas_cabecalho: list[str],
    secoes: list[tuple[str, list[str], list[tuple[str, ...]]]],
) -> Path:
    """Monta o documento e salva em `caminho`. Cria o diretório pai se não
    existir. Times New Roman no documento inteiro, título (nível 0) em preto,
    subtítulo de seção (nível 1) em dourado.

    `secoes`: lista de `(subtitulo, colunas, linhas)` — `colunas` é o
    cabeçalho da tabela (N textos), `linhas` é lista de tuplas com o mesmo N
    de valores, uma por item. Seção sem linha nenhuma vira só um parágrafo
    "Nenhum item nesta seção.", sem tabela vazia.
    """
    document = Document()
    _forcar_fonte(document.styles["Normal"])
    _forcar_fonte(document.styles["Title"], cor=_PRETO)
    _forcar_fonte(document.styles["Heading 1"], cor=_DOURADO)

    document.add_heading(_limpar(titulo), 0)
    for linha in linhas_cabecalho:
        document.add_paragraph(_limpar(linha))

    for subtitulo, colunas, linhas in secoes:
        document.add_heading(_limpar(subtitulo), level=1)
        if not linhas:
            document.add_paragraph("Nenhum item nesta seção.")
            continue
        tabela = document.add_table(rows=1, cols=len(colunas))
        tabela.style = "Table Grid"
        for celula, texto in zip(tabela.rows[0].cells, colunas):
            celula.text = _limpar(texto)
        for linha in linhas:
            celulas = tabela.add_row().cells
            for celula, valor in zip(celulas, linha):
                celula.text = _limpar(valor)

    try:
        caminho.parent.mkdir(parents=True, exist_ok=True)
        document.save(caminho)
    except OSError as exc:
        raise DocxReportError(f"falha ao salvar '{caminho}'") from exc
    return caminho


def gerar_com_seguranca(
    gerador: Callable[[], Path | None], contexto: str,
) -> tuple[Path | None, bool]:
    """Roda `gerador` sem deixar `DocxReportError`/`ValueError` propagar — relatório é
    secundário ao resultado principal, já gravado antes de chegar aqui.

    Devolve `(caminho, falhou)`: `(Path, False)` gerou, `(None, False)` o
    gerador decidiu não gerar nada, `(None, True)` a geração falhou.
    `contexto` só identifica o item no log (ex. "divergencia nota 123").
    """
    try:
        return gerador(), False
    except (DocxReportError, ValueError):
        log.exception("docx_report: falha gerando relatorio .docx (%s)", contexto)
        return None, True
